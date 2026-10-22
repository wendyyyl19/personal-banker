# Personal Banker 
# This is my final version of the Personal Banker

# I use openpyxl is for Excel files, os is for the file location and date is for dates
#https://www.geeksforgeeks.org/python/introduction-to-python-openpyxl/
from openpyxl import Workbook, load_workbook
import os
from datetime import date

# lists and append from Week 2; lists and dictionaries from Week 7
expenses = []  # all the expense records

# References:
#https://www.geeksforgeeks.org/python/python-os-path-abspath-method-with-example/
#https://docs.python.org/3/library/os.path.html
# abspath gets the full script path; dirname gets its folder; join adds the file name.
# The Excel file is stored in the same folder as this Python file.
# The reflection notes the old Desktop path would not work on every computer.
folder = os.path.dirname(os.path.abspath(__file__))
filename = os.path.join(folder, "expenses.xlsx")


# functions and parameters from Week 2
# input, while True, if/elif/else and break
def addexpense():
    print("Choose a category:")
    print("1. Food")
    print("2. Transport")
    print("3. Study")
    print("4. Other")

    while True:
        choice = input("Choose 1 to 4: ")
        if choice in ["1", "2", "3", "4"]:
            break
        print("Please choose 1, 2, 3 or 4.")

    # use if/elif to choose the category
    # categories = {1: "Food", 2: "Transport", 3: "Study", 4: "Other"}
    # category = categories[int(choice)]
    if choice == "1":
        category = "Food"
    elif choice == "2":
        category = "Transport"
    elif choice == "3":
        category = "Study"
    else:
        category = "Other"

    # float() turns the text into a number.
    #https://docs.python.org/3/library/functions.html#float
    # Non-numeric input will cause an error, I know that is still there.
    while True:
        money = float(input("Amount spent in AUD: $"))  # turn the input into a number
        if money > 0:
            break
        print("Please enter an amount greater than zero.")

    # I tried date_text = input("Date: ") at first and typed 21/10/2026,
    # then split("-") broke it, so I now ask for YYYY-MM-DD.
    # https://docs.python.org/3/library/stdtypes.html#str.split
    mydate = input("Date (YYYY-MM-DD): ")
    y, m, d = mydate.split("-")  # y is year, m is month, d is day
    dayspent = date(int(y), int(m), int(d))

    # keep each expense in the list
    record = {
        "date": dayspent,
        "category": category,
        "amount": money
    }

    expenses.append(record)  # add the new record to the list


# total = 0, then for item in [1, 3, 5]: total = total + item, from Week 7.
# return ends this function when the list is empty; it does not exit the menu.
#https://docs.python.org/3/tutorial/controlflow.html#defining-functions
def showexpenses():
    if len(expenses) == 0:
        print("No expenses yet.")
        return

    total = 0  # keep this outside the loop so it does not reset for each record

    print("My expenses are here:")
    # The reflection replay showed expenses[0] only displayed the first record.
    for e in expenses:
        print("Date:", e["date"])
        print("Category:", e["category"])
        print("Amount: $", e["amount"])
        total = total + e["amount"]  # add this expense to the total

    print("Total spent: $", total)


# I got the weekly total wrong at first. I thought records were missing, but actually I was only comparing week numbers, and week 1 of 2026 and week 1 of 2025 have the same number. Adding the year check fixed it.
#https://docs.python.org/3/library/datetime.html#datetime.date.isocalendar
def showmyweek():
    today = date.today()
    thisyear, thisweek, weekday = today.isocalendar()  # weekday is not used here
    weektotal = 0

    for e in expenses:
        year, week, weekday = e["date"].isocalendar()
        # compare the year too, because week numbers repeat each year
        if year == thisyear and week == thisweek:
            weektotal = weektotal + e["amount"]

    # .2f keeps two decimal places
    print(f"This week's total: ${weektotal:.2f}")
    # I use 1 AUD = 5 RMB for a rough estimate.
    # My mom suggested using a live rate in the future, but I have not added it yet.
    print(f"Rough estimate: RMB {weektotal * 5:.2f} (using 1 AUD = 5 RMB)")


# totals per category, from Week 7 dict notes
def showmycategories():
    # Other can group different purchases together, so the total hides their purpose.
    totals = {"Food": 0, "Transport": 0, "Study": 0, "Other": 0}
    for e in expenses:
        category = e["category"]
        totals[category] = totals[category] + e["amount"]

    print("Category totals for all saved expenses:")
    for category in totals:
        print(category, f"${totals[category]:.2f}")


# save to Excel
def saveexcel():
    book = Workbook()
    sheet = book.active  # use the first sheet
    sheet.title = "Weekly Expenses"
    sheet.append(["Date", "Category", "Amount (AUD)", "Amount (RMB estimate, 1 AUD = 5 RMB)"])

    # Save the previous records together with the new records
    for e in expenses:
        sheet.append([
            e["date"].strftime("%Y-%m-%d"),
            e["category"],
            e["amount"],
            round(e["amount"] * 5, 2)
        ])

    book.save(filename)
    book.close()
    print("Saved to:", filename)


# Excel reference: https://openpyxl.readthedocs.io/en/stable/tutorial.html
# See Loading from a file and iter_rows(values_only=True).
# Row 1 has headings, so the expense records start at row 2
# The reflection separates saving from reopening: saved rows still need loading.
def loadexcel():
    if not os.path.exists(filename):
        print("No saved file yet. Start with a new list.")
        return

    book = load_workbook(filename)
    sheet = book["Weekly Expenses"]

    #https://openpyxl.readthedocs.io/en/stable/tutorial.html#accessing-many-cells
    for row in sheet.iter_rows(min_row=2, values_only=True):  # skip headings, read values only
        mydate = row[0]  # first column is the date
        category = row[1]  # second column is the category
        money = row[2]  # third column is the AUD amount
        if mydate is not None:
            y, m, d = mydate.split("-")  # y is year, m is month, d is day
            dayspent = date(int(y), int(m), int(d))
            expenses.append({"date": dayspent, "category": category, "amount": money})

    book.close()
    print("Loaded records:", len(expenses))


# Load the saved records before showing the menu
loadexcel()

# conditionals_extra.py ticket example repeats input and uses break when the user enters x. This menu applies that flow to numbered choices.
# Keep showing this menu until option 5 is chosen
while True:
    print("\nPersonal Banker")
    print("1. Add expense")
    print("2. Show expenses and totals")
    print("3. Show category totals")
    print("4. Save to Excel")
    print("5. Save and exit")

    choice = input("Choose 1 to 5: ")
    if choice == "1":
        addexpense()
        # Save after adding an expense so it is available next time
        saveexcel()
    elif choice == "2":
        showexpenses()
        showmyweek()
    elif choice == "3":
        showmycategories()
    elif choice == "4":
        saveexcel()
    elif choice == "5":
        saveexcel()
        print("Goodbye!")
        break
    else:
        print("Please choose 1, 2, 3, 4 or 5.")